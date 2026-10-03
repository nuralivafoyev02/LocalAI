import json

import pandas as pd
import pytest

from localai.calculator import calculate
from localai.permissions import GrantStore
from localai.tools import TOOLS_BY_NAME, ToolContext, ToolError, find_tool, needs_approval


def run(settings, tool_name, **args):
    tool = TOOLS_BY_NAME[tool_name]
    ctx = ToolContext(settings=settings, chat_id="t")
    prepared = tool.prepare(ctx, args)
    return tool.run(ctx, prepared)


def make_pdf(path, text):
    """Matnli eng kichik PDF fayl yaratadi."""
    stream = f"BT /F1 18 Tf 72 720 Td ({text}) Tj ET".encode()
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for number, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += b"%d 0 obj\n" % number + body + b"\nendobj\n"
    xref = len(out)
    out += b"xref\n0 %d\n0000000000 65535 f \n" % (len(objects) + 1)
    for offset in offsets:
        out += b"%010d 00000 n \n" % offset
    out += b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (len(objects) + 1, xref)
    path.write_bytes(bytes(out))


@pytest.fixture
def project(workspace):
    root = workspace / "loyiha"
    (root / "src").mkdir(parents=True)
    (root / "node_modules" / "pkg").mkdir(parents=True)
    (root / "README.md").write_text("# Loyiha\n\nBu sinov loyihasi.\n", encoding="utf-8")
    (root / "src" / "main.py").write_text("\n".join(f"line {i}" for i in range(1, 1001)), encoding="utf-8")
    (root / "src" / "util.py").write_text("def yordamchi():\n    return 'TOPILDI'\n", encoding="utf-8")
    (root / ".hidden").write_text("x")
    (root / "image.bin").write_bytes(b"\x00\x01\x02" * 100)
    return root


def test_list_directory(settings, project):
    result = run(settings, "list_directory", path=str(project), depth=2)
    assert "README.md" in result.content
    assert "src/" in result.content and "main.py" in result.content
    assert "node_modules/  (o'tkazib yuborildi)" in result.content
    assert "pkg" not in result.content
    assert ".hidden" not in result.content and "yashirin" in result.content
    assert "papka" in result.summary


def test_list_directory_rejects_file(settings, project):
    with pytest.raises(ToolError):
        run(settings, "list_directory", path=str(project / "README.md"))


def test_read_file_with_paging(settings, project):
    result = run(settings, "read_file", path=str(project / "src" / "main.py"), start_line=10, max_lines=5)
    assert "10│ line 10" in result.content and "14│ line 14" in result.content
    assert "line 15" not in result.content
    assert "start_line=15" in result.content


def test_read_file_respects_output_budget(settings, project):
    settings.tool_output_chars = 2000
    result = run(settings, "read_file", path=str(project / "src" / "main.py"))
    assert len(result.content) < 2600
    assert "Davomi bor" in result.content


def test_read_binary_file_fails(settings, project):
    with pytest.raises(ToolError, match="binary"):
        run(settings, "read_file", path=str(project / "image.bin"))


def test_read_missing_file(settings, project):
    with pytest.raises(ToolError, match="topilmadi"):
        run(settings, "read_file", path=str(project / "yoq.txt"))


def test_read_cp1251_text(settings, workspace):
    path = workspace / "ru.txt"
    path.write_bytes("Привет, мир".encode("cp1251"))
    assert "Привет, мир" in run(settings, "read_file", path=str(path)).content


def test_read_docx(settings, workspace):
    from docx import Document

    document = Document()
    document.add_paragraph("Salom, hujjat!")
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "A"
    table.rows[0].cells[1].text = "B"
    document.save(workspace / "hujjat.docx")
    result = run(settings, "read_file", path=str(workspace / "hujjat.docx"))
    assert "Salom, hujjat!" in result.content and "A | B" in result.content


def test_read_pdf(settings, workspace):
    make_pdf(workspace / "a.pdf", "Salom PDF")
    result = run(settings, "read_file", path=str(workspace / "a.pdf"))
    assert "Salom PDF" in result.content and "1-sahifa" in result.content


def test_read_typo_path_is_corrected(settings, project):
    result = run(settings, "read_file", path=str(project / "REDME.md"))
    assert "Bu sinov loyihasi" in result.content
    assert "tuzatildi" in result.content


def test_search_files(settings, project):
    by_name = run(settings, "search_files", path=str(project), name="*.py")
    assert "main.py" in by_name.content and "util.py" in by_name.content
    by_text = run(settings, "search_files", path=str(project), text="topildi")
    assert "util.py" in by_text.content and "2:" in by_text.content
    missing = run(settings, "search_files", path=str(project), name="yoq_narsa")
    assert missing.summary == "topilmadi"
    with pytest.raises(ToolError):
        run(settings, "search_files", path=str(project))


def test_search_skips_sensitive_files(settings, project):
    (project / ".env").write_text("SECRET=topildi")
    (project / "keys").mkdir()
    (project / "keys" / "server.pem").write_text("topildi")
    result = run(settings, "search_files", path=str(project), text="topildi")
    assert ".env" not in result.content and "server.pem" not in result.content


@pytest.fixture
def sales(workspace):
    frame = pd.DataFrame({
        "Mahsulot": ["Olma", "Nok", "Olma", "Uzum", "Nok", "Olma"],
        "Viloyat": ["Toshkent", "Samarqand", "Samarqand", "Toshkent", "Toshkent", "Buxoro"],
        "Summa": [100, 250, 150, 400, 50, 300],
    })
    path = workspace / "savdo.xlsx"
    with pd.ExcelWriter(path) as writer:
        frame.to_excel(writer, sheet_name="Savdo", index=False)
        pd.DataFrame({"x": [1]}).to_excel(writer, sheet_name="Boshqa", index=False)
    frame.to_csv(workspace / "savdo.csv", index=False, sep=";")
    return path


def test_table_overview(settings, sales):
    result = run(settings, "analyze_table", path=str(sales))
    assert "Varaqlar (2): Savdo, Boshqa" in result.content
    assert "Summa" in result.content and "1250" in result.content


def test_table_group_by_with_fuzzy_names(settings, sales):
    result = run(
        settings, "analyze_table", path=str(sales), sheet="savdo", group_by=["mahsulot"],
        aggregate=[{"column": "sum", "func": "sum"}, {"column": "*", "func": "count"}],
        sort_by="Summa_sum", descending=True,
    )
    lines = [line for line in result.content.splitlines() if line.startswith("| ")]
    assert lines[0].startswith("| Mahsulot")
    assert "| Olma | 550 | 3 |" in lines[2]
    assert "| Uzum | 400 | 1 |" in lines[3]


def test_table_filters(settings, sales):
    result = run(
        settings, "analyze_table", path=str(sales.with_suffix(".csv")),
        filters=[{"column": "Viloyat", "op": "=", "value": "toshkent"}, {"column": "Summa", "op": ">", "value": "60"}],
    )
    assert "Filtrdan keyin: 2 qator" in result.content
    assert "Uzum" in result.content and "Nok" not in result.content


def test_table_dict_style_arguments(settings, sales):
    result = run(settings, "analyze_table", path=str(sales), filters={"Mahsulot": "Olma"}, aggregate={"Summa": "o'rtacha"})
    assert "Summa_mean" in result.content and "183.333" in result.content


def test_table_unknown_column(settings, sales):
    with pytest.raises(ToolError, match="Mavjud ustunlar"):
        run(settings, "analyze_table", path=str(sales), group_by=["qwerty_yoq"])


def test_table_json_records(settings, workspace):
    (workspace / "data.json").write_text(json.dumps({"items": [{"a": 1, "b": {"c": 2}}, {"a": 3, "b": {"c": 4}}]}))
    result = run(settings, "analyze_table", path=str(workspace / "data.json"), aggregate=[{"column": "b.c", "func": "sum"}])
    assert "| 6 |" in result.content


def test_table_rejects_other_formats(settings, project):
    with pytest.raises(ToolError):
        run(settings, "analyze_table", path=str(project / "README.md"))


def test_write_file_creates_and_overwrites(settings, workspace):
    target = workspace / "yangi" / "skript.py"
    result = run(settings, "write_file", path=str(target), content="print(1)\n")
    assert target.read_text() == "print(1)\n" and "yaratildi" in result.summary
    result = run(settings, "write_file", path=str(target), content="print(2)\n")
    assert target.read_text() == "print(2)\n" and "yangilandi" in result.summary
    assert not list(target.parent.glob(".*localai-tmp"))


def test_write_file_requires_content(settings, workspace):
    with pytest.raises(ToolError):
        run(settings, "write_file", path=str(workspace / "a.txt"))


@pytest.mark.parametrize("expression, expected", [
    ("2+2*2", "6"),
    ("(1250000 * 12) * 0.88", "13 200 000"),
    ("12 500 000 * 12%", "1 500 000"),
    ("sqrt(16) + 2^3", "12"),
    ("15% dan 200", "30"),
    ("round(10/3, 2)", "3.33"),
    ("factorial(5)", "120"),
    ("mean(1, 2, 3, 4)", "2.5"),
    ("10 ÷ 4", "2.5"),
])
def test_calculate(expression, expected):
    assert calculate(expression) == expected


@pytest.mark.parametrize("expression", ["__import__('os')", "9**9**9", "1/0", "open('x')", "a.b", "factorial(10**6)"])
def test_calculate_rejects_unsafe(expression):
    with pytest.raises(ValueError):
        calculate(expression)


def test_find_tool_is_typo_tolerant():
    assert find_tool("read_file").name == "read_file"
    assert find_tool("read_files").name == "read_file"
    assert find_tool("ls").name == "list_directory"
    assert find_tool("calculator").name == "calculate"
    assert find_tool("delete_everything") is None


def test_needs_approval(settings, project):
    grants = GrantStore()
    ctx = ToolContext(settings=settings, chat_id="c")
    read = TOOLS_BY_NAME["read_file"]
    prepared = read.prepare(ctx, {"path": str(project / "README.md")})
    assert needs_approval(read, prepared, ctx, grants)
    grants.grant("c", project)
    assert not needs_approval(read, prepared, ctx, grants)

    write = TOOLS_BY_NAME["write_file"]
    prepared = write.prepare(ctx, {"path": str(project / "x.txt"), "content": "x"})
    assert needs_approval(write, prepared, ctx, grants)

    calc = TOOLS_BY_NAME["calculate"]
    assert not needs_approval(calc, calc.prepare(ctx, {"expression": "1+1"}), ctx, grants)

    upload = settings.upload_dir / "abc12345" / "f.txt"
    upload.parent.mkdir(parents=True)
    upload.write_text("x")
    prepared = read.prepare(ctx, {"path": str(upload)})
    assert not needs_approval(read, prepared, ctx, GrantStore())
