"""Fayllardan matn va jadval ajratib olish (matn, kod, PDF, DOCX, Excel, CSV, JSON)."""

import difflib
import json
import threading
from collections import OrderedDict
from pathlib import Path

import pandas as pd


TABLE_SUFFIXES = {".csv", ".tsv", ".xlsx", ".xlsm", ".xls"}
DOCUMENT_SUFFIXES = {".pdf", ".docx"}
MAX_TEXT_BYTES = 5 * 1024 * 1024
MAX_DOCUMENT_BYTES = 50 * 1024 * 1024
SKIP_DIRS = {
    ".git", ".hg", ".svn", "node_modules", "__pycache__", ".venv", "venv", "env",
    ".mypy_cache", ".pytest_cache", ".ruff_cache", ".tox", ".idea", ".vscode",
    ".next", ".nuxt", "dist", "build", "target", ".gradle", ".cache", "Pods",
    ".DS_Store", "site-packages", ".terraform",
}


class FileReadError(ValueError):
    pass


def human_size(size):
    value = float(size)
    for unit in ("B", "KB", "MB", "GB"):
        if value < 1024 or unit == "GB":
            return f"{value:.0f} {unit}" if unit == "B" else f"{value:.1f} {unit}"
        value /= 1024
    return f"{value:.1f} GB"


def is_binary(sample):
    if b"\x00" in sample:
        return True
    if not sample:
        return False
    control = sum(1 for byte in sample if byte < 9 or 13 < byte < 32)
    return control / len(sample) > 0.3


def decode_bytes(data):
    for encoding in ("utf-8-sig", "cp1251"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("latin-1")


def read_text(path):
    path = Path(path)
    size = path.stat().st_size
    if size > MAX_TEXT_BYTES:
        raise FileReadError(
            f"Fayl juda katta ({human_size(size)}). Matnli fayllar uchun limit {human_size(MAX_TEXT_BYTES)}."
        )
    data = path.read_bytes()
    if is_binary(data[:8192]):
        raise FileReadError("Bu ikkilik (binary) fayl, uni matn sifatida o'qib bo'lmaydi.")
    return decode_bytes(data)


def read_pdf(path):
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    if reader.is_encrypted:
        try:
            reader.decrypt("")
        except Exception as error:
            raise FileReadError("PDF parol bilan himoyalangan.") from error
    pages = []
    for number, page in enumerate(reader.pages, start=1):
        text = (page.extract_text() or "").strip()
        pages.append(f"--- {number}-sahifa ---\n{text}" if text else f"--- {number}-sahifa (matn yo'q) ---")
    text = "\n\n".join(pages)
    if not any(page.extract_text() for page in reader.pages[:3]) and len(reader.pages) > 0:
        text += "\n\n[Eslatma: PDF skanerlangan rasm bo'lishi mumkin, matn ajratilmadi.]"
    return text


def read_docx(path):
    from docx import Document

    document = Document(str(path))
    parts = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
    for index, table in enumerate(document.tables, start=1):
        rows = []
        for row in table.rows:
            rows.append(" | ".join(cell.text.strip() for cell in row.cells))
        if rows:
            parts.append(f"\n[Jadval {index}]\n" + "\n".join(rows))
    return "\n".join(parts)


def read_notebook(path):
    notebook = json.loads(read_text(path))
    cells = []
    for index, cell in enumerate(notebook.get("cells", []), start=1):
        source = cell.get("source", "")
        if isinstance(source, list):
            source = "".join(source)
        cells.append(f"# [{index}] {cell.get('cell_type', 'cell')}\n{source}")
    return "\n\n".join(cells)


def extract_text(path):
    """Faylning matnli ko'rinishini qaytaradi: (tur, matn)."""
    path = Path(path)
    suffix = path.suffix.lower()
    size = path.stat().st_size
    if suffix in DOCUMENT_SUFFIXES or suffix in TABLE_SUFFIXES:
        if size > MAX_DOCUMENT_BYTES:
            raise FileReadError(f"Fayl juda katta ({human_size(size)}).")
    try:
        if suffix == ".pdf":
            return "pdf", read_pdf(path)
        if suffix == ".docx":
            return "docx", read_docx(path)
        if suffix == ".ipynb":
            return "notebook", read_notebook(path)
        if suffix in TABLE_SUFFIXES:
            frames = load_tables(path)
            return "table", table_overview(frames, path.name)
    except FileReadError:
        raise
    except ImportError as error:
        raise FileReadError("Kerakli kutubxona o'rnatilmagan: " + str(error)) from error
    except Exception as error:
        raise FileReadError(f"Faylni o'qib bo'lmadi: {error}") from error
    if suffix in {".doc", ".ppt", ".pptx", ".odt", ".rtf", ".pages", ".key", ".numbers"}:
        raise FileReadError(f"{suffix} formati hozircha qo'llanmaydi. Faylni PDF, DOCX yoki TXT ko'rinishida saqlang.")
    return "text", read_text(path)


# ----------------------------------------------------------------- jadvallar

_table_cache = OrderedDict()
_cache_lock = threading.Lock()


def _json_frame(value):
    if isinstance(value, list):
        if value and all(isinstance(row, dict) for row in value):
            return pd.json_normalize(value, sep=".")
        return pd.DataFrame({"value": value})
    if isinstance(value, dict):
        list_values = [item for item in value.values() if isinstance(item, list)]
        if len(list_values) == 1:
            records = list_values[0]
            if records and all(isinstance(row, dict) for row in records):
                return pd.json_normalize(records, sep=".")
            if records:
                return pd.DataFrame({"value": records})
        return pd.json_normalize(value, sep=".")
    return pd.DataFrame({"value": [value]})


def _load_tables_uncached(path):
    suffix = path.suffix.lower()
    if suffix in {".xlsx", ".xlsm", ".xls"}:
        try:
            return pd.read_excel(path, sheet_name=None)
        except ImportError:
            raise
        except Exception as error:
            raise FileReadError(f"Excel faylini o'qib bo'lmadi: fayl buzilgan yoki boshqa formatda ({error}).") from error
    if suffix == ".json":
        return {"JSON": _json_frame(json.loads(read_text(path)))}
    if suffix == ".tsv":
        return {path.stem: pd.read_csv(path, sep="\t", encoding_errors="replace")}
    if suffix == ".csv":
        text = read_text(path)
        from io import StringIO

        try:
            frame = pd.read_csv(StringIO(text), sep=None, engine="python")
        except Exception:
            frame = pd.read_csv(StringIO(text))
        return {path.stem: frame}
    raise FileReadError("Jadval formati qo'llanmaydi. CSV, TSV, XLSX, XLS yoki JSON kerak.")


def load_tables(path):
    path = Path(path)
    stat = path.stat()
    if stat.st_size > MAX_DOCUMENT_BYTES:
        raise FileReadError(f"Jadval fayli juda katta ({human_size(stat.st_size)}).")
    key = (str(path), stat.st_mtime_ns, stat.st_size)
    with _cache_lock:
        if key in _table_cache:
            _table_cache.move_to_end(key)
            return _table_cache[key]
    frames = _load_tables_uncached(path)
    frames = {str(name): frame for name, frame in frames.items()}
    with _cache_lock:
        _table_cache[key] = frames
        while len(_table_cache) > 6:
            _table_cache.popitem(last=False)
    return frames


def _cell(value):
    if value is None:
        return ""
    try:
        if pd.isna(value):
            return ""
    except (TypeError, ValueError):
        pass
    if isinstance(value, float):
        text = f"{value:.6g}" if abs(value) < 1e15 else f"{value:.4e}"
    else:
        text = str(value)
    text = text.replace("|", "\\|").replace("\n", " ").strip()
    return text[:80] + "…" if len(text) > 80 else text


def frame_to_markdown(frame, max_rows=50):
    if frame.empty:
        return "(bo'sh natija)"
    shown = frame.head(max_rows)
    columns = [str(column) for column in shown.columns]
    lines = [
        "| " + " | ".join(_cell(column) for column in columns) + " |",
        "| " + " | ".join("---" for _ in columns) + " |",
    ]
    for row in shown.itertuples(index=False, name=None):
        lines.append("| " + " | ".join(_cell(value) for value in row) + " |")
    if len(frame) > max_rows:
        lines.append(f"\n… yana {len(frame) - max_rows} qator ko'rsatilmadi.")
    return "\n".join(lines)


def match_column(frame, name):
    """Ustun nomini xatolarga chidamli tarzda topadi."""
    columns = [str(column) for column in frame.columns]
    name = str(name).strip()
    if name in columns:
        return frame.columns[columns.index(name)]
    lowered = [column.lower() for column in columns]
    if name.lower() in lowered:
        return frame.columns[lowered.index(name.lower())]
    match = difflib.get_close_matches(name.lower(), lowered, n=1, cutoff=0.6)
    if match:
        return frame.columns[lowered.index(match[0])]
    raise FileReadError(f"'{name}' ustuni topilmadi. Mavjud ustunlar: {', '.join(columns)}")


def pick_sheet(frames, sheet):
    if not frames:
        raise FileReadError("Faylda jadval topilmadi.")
    names = list(frames)
    if sheet in (None, ""):
        return names[0]
    sheet = str(sheet)
    if sheet in frames:
        return sheet
    lowered = [name.lower() for name in names]
    match = difflib.get_close_matches(sheet.lower(), lowered, n=1, cutoff=0.5)
    if match:
        return names[lowered.index(match[0])]
    if sheet.isdigit() and 0 < int(sheet) <= len(names):
        return names[int(sheet) - 1]
    raise FileReadError(f"'{sheet}' varag'i topilmadi. Varaqlar: {', '.join(names)}")


def describe_frame(frame, preview_rows=15):
    parts = [f"O'lcham: {len(frame):,} qator × {len(frame.columns)} ustun"]
    column_lines = []
    for column in frame.columns:
        series = frame[column]
        missing = int(series.isna().sum())
        info = f"- {column} ({series.dtype}"
        info += f", {missing} bo'sh)" if missing else ")"
        column_lines.append(info)
    parts.append("Ustunlar:\n" + "\n".join(column_lines))

    numeric = frame.select_dtypes(include="number")
    if not numeric.empty:
        stats = numeric.describe().T
        stats.insert(0, "sum", numeric.sum())
        stats = stats.reset_index().rename(columns={"index": "ustun"})
        parts.append("Raqamli ustunlar statistikasi (butun fayl bo'yicha):\n" + frame_to_markdown(stats.round(4), 40))

    tops = []
    for column in [column for column in frame.columns if column not in numeric.columns]:
        if len(tops) >= 8:
            break
        text = frame[column].dropna().astype(str)
        unique = text.nunique()
        if not unique or unique > max(50, len(frame) // 2):
            continue
        counts = text.value_counts().head(5)
        values = ", ".join(f"{_cell(index)} ({count})" for index, count in counts.items())
        tops.append(f"- {column}: {unique} xil qiymat; ko'p uchraydi: {values}")
    if tops:
        parts.append("Matnli ustunlardagi eng ko'p qiymatlar:\n" + "\n".join(tops))

    if preview_rows:
        parts.append(f"Dastlabki {min(preview_rows, len(frame))} qator:\n" + frame_to_markdown(frame, preview_rows))
    return "\n\n".join(parts)


def table_overview(frames, filename, sheet=None, preview_rows=15):
    names = list(frames)
    header = f"Fayl: {filename}"
    if len(names) > 1:
        header += f"\nVaraqlar ({len(names)}): " + ", ".join(names)
    selected = [pick_sheet(frames, sheet)] if sheet or len(names) == 1 else names[:3]
    sections = [header]
    for name in selected:
        title = f"## Varaq: {name}" if len(names) > 1 or sheet else ""
        body = describe_frame(frames[name], preview_rows if len(selected) == 1 else 8)
        sections.append((title + "\n" if title else "") + body)
    if len(names) > len(selected):
        sections.append(f"(Qolgan {len(names) - len(selected)} varaqni ko'rish uchun 'sheet' parametrini bering.)")
    return "\n\n".join(sections)


FILTER_OPS = {
    "==", "!=", ">", ">=", "<", "<=", "contains", "not_contains", "startswith",
    "endswith", "in", "not_in", "isnull", "notnull",
}
OP_ALIASES = {
    "=": "==", "eq": "==", "equals": "==", "is": "==", "ne": "!=", "<>": "!=", "not": "!=",
    "gt": ">", "gte": ">=", "ge": ">=", "lt": "<", "lte": "<=", "le": "<=",
    "like": "contains", "has": "contains", "includes": "contains", "starts_with": "startswith",
    "ends_with": "endswith", "empty": "isnull", "is_null": "isnull", "not_empty": "notnull",
    "not_null": "notnull",
}
AGG_FUNCS = {"sum", "mean", "median", "min", "max", "count", "nunique", "std", "first", "last"}
AGG_ALIASES = {
    "avg": "mean", "average": "mean", "o'rtacha": "mean", "ortacha": "mean", "total": "sum",
    "yig'indi": "sum", "jami": "sum", "soni": "count", "unique": "nunique", "distinct": "nunique",
}


def _as_list(value):
    if value is None or value == "":
        return []
    if isinstance(value, str):
        return [item.strip() for item in value.split(",") if item.strip()]
    if isinstance(value, (list, tuple)):
        return list(value)
    return [value]


def normalize_filters(filters):
    if not filters:
        return []
    if isinstance(filters, str):
        try:
            filters = json.loads(filters)
        except ValueError:
            raise FileReadError("filters JSON ko'rinishida bo'lishi kerak.")
    if isinstance(filters, dict):
        if "column" in filters:
            filters = [filters]
        else:
            filters = [{"column": key, "op": "==", "value": value} for key, value in filters.items()]
    normalized = []
    for item in filters:
        if not isinstance(item, dict) or "column" not in item:
            raise FileReadError("Har bir filtr {column, op, value} ko'rinishida bo'lishi kerak.")
        op = str(item.get("op", "==")).strip().lower()
        op = OP_ALIASES.get(op, op)
        if op not in FILTER_OPS:
            raise FileReadError(f"Noma'lum filtr operatori: {op}. Mavjud: {', '.join(sorted(FILTER_OPS))}")
        normalized.append({"column": item["column"], "op": op, "value": item.get("value")})
    return normalized


def normalize_aggregate(aggregate):
    if not aggregate:
        return []
    if isinstance(aggregate, str):
        try:
            aggregate = json.loads(aggregate)
        except ValueError:
            aggregate = [{"column": "*", "func": aggregate}]
    if isinstance(aggregate, dict):
        if "func" in aggregate:
            aggregate = [aggregate]
        else:
            aggregate = [{"column": key, "func": value} for key, value in aggregate.items()]
    normalized = []
    for item in aggregate:
        if isinstance(item, str):
            item = {"column": "*", "func": item}
        func = str(item.get("func", "sum")).strip().lower()
        func = AGG_ALIASES.get(func, func)
        if func not in AGG_FUNCS:
            raise FileReadError(f"Noma'lum agregat funksiya: {func}. Mavjud: {', '.join(sorted(AGG_FUNCS))}")
        normalized.append({"column": item.get("column") or "*", "func": func})
    return normalized


def _numeric(value):
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return value
    try:
        return float(str(value).replace(" ", "").replace(",", "."))
    except (TypeError, ValueError):
        return None


def _apply_filter(frame, spec):
    column = match_column(frame, spec["column"])
    series = frame[column]
    op, value = spec["op"], spec["value"]
    if op == "isnull":
        return series.isna()
    if op == "notnull":
        return series.notna()
    if op in {"in", "not_in"}:
        values = [str(item).lower() for item in _as_list(value)]
        mask = series.astype(str).str.lower().isin(values)
        return ~mask if op == "not_in" else mask
    text = series.astype(str).str.lower()
    needle = str(value).lower()
    if op == "contains":
        return text.str.contains(needle, regex=False, na=False)
    if op == "not_contains":
        return ~text.str.contains(needle, regex=False, na=False)
    if op == "startswith":
        return text.str.startswith(needle, na=False)
    if op == "endswith":
        return text.str.endswith(needle, na=False)
    number = _numeric(value)
    if pd.api.types.is_datetime64_any_dtype(series):
        left, right = series, pd.to_datetime(str(value), errors="coerce")
    elif number is not None and (pd.api.types.is_numeric_dtype(series) or op in {">", ">=", "<", "<="}):
        left = pd.to_numeric(series, errors="coerce")
        right = number
    else:
        left, right = text, needle
    if op == "==":
        return left == right
    if op == "!=":
        return left != right
    if op == ">":
        return left > right
    if op == ">=":
        return left >= right
    if op == "<":
        return left < right
    return left <= right


def query_table(frame, columns=None, filters=None, group_by=None, aggregate=None,
                sort_by=None, descending=False, limit=50):
    notes = []
    result = frame
    for spec in normalize_filters(filters):
        result = result[_apply_filter(result, spec)]
    if filters:
        notes.append(f"Filtrdan keyin: {len(result):,} qator (jami {len(frame):,}).")

    group_columns = [match_column(result, name) for name in _as_list(group_by)]
    aggregations = normalize_aggregate(aggregate)

    if group_columns or aggregations:
        if not aggregations:
            aggregations = [{"column": "*", "func": "count"}]
        named = {}
        for spec in aggregations:
            if spec["column"] in ("*", "", None):
                named["soni"] = None
                continue
            column = match_column(result, spec["column"])
            if spec["func"] in {"sum", "mean", "median", "std"} and not pd.api.types.is_numeric_dtype(result[column]):
                result = result.assign(**{str(column): pd.to_numeric(result[column], errors="coerce")})
            named[f"{column}_{spec['func']}"] = (column, spec["func"])
        if group_columns:
            grouped = result.groupby(group_columns, dropna=False)
            pieces = {
                label: grouped.size() if spec is None else grouped[spec[0]].agg(spec[1])
                for label, spec in named.items()
            }
            result = pd.DataFrame(pieces).reset_index()
        else:
            result = pd.DataFrame([{
                label: len(result) if spec is None else result[spec[0]].agg(spec[1])
                for label, spec in named.items()
            }])

    if columns and not (group_columns or aggregations):
        result = result[[match_column(result, name) for name in _as_list(columns)]]

    if sort_by:
        try:
            sort_column = match_column(result, sort_by)
            result = result.sort_values(sort_column, ascending=not descending, kind="stable")
        except FileReadError:
            notes.append(f"'{sort_by}' bo'yicha saralab bo'lmadi.")

    try:
        limit = max(1, min(int(limit or 50), 200))
    except (TypeError, ValueError):
        limit = 50
    numeric_cols = result.select_dtypes(include="number").columns
    if len(numeric_cols):
        result = result.copy()
        result[numeric_cols] = result[numeric_cols].round(4)
    return result, limit, notes
